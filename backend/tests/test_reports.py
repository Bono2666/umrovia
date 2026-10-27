"""S17: laporan siap pakai — rekonsiliasi terhadap catatan sumber (RPT01), pemisahan wallet vs kas (RPT02), ekspor, izin."""
from datetime import timedelta
from io import BytesIO
from uuid import uuid4
from zoneinfo import ZoneInfo

import pytest
from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import connection
from django.db.models import Sum
from django.utils import timezone
from openpyxl import load_workbook

from modules.audit.models import AuditEvent
from modules.payments.models import TravelInvoice, TravelPayment
from modules.reports import export
from modules.travel_finance.models import TravelExpense
from modules.wallet.models import WalletLedger

from .test_billing import billing as billing
from .test_billing import fund
from .test_payment_security import staff
from .test_payments import ok, post
from .test_payments import scenario as scenario
from .test_wallet import wallet_setup as wallet_setup

TODAY = timezone.now().astimezone(ZoneInfo("Asia/Jakarta")).date()  # tanggal lokal travel, bukan UTC
WIDE = f"?from={(TODAY - timedelta(days=5)).isoformat()}&to={(TODAY + timedelta(days=90)).isoformat()}"


def kpis(report):
    return {kpi["key"]: kpi["value"] for kpi in report["kpis"]}


def section(report, key):
    return next(s for s in report["sections"] if s["key"] == key)


def pay(travel, index, amount, *, method="CASH", bank=None, client=None):
    line = travel["invoice"]["lines"][index]
    body = {
        "invoice": travel["invoice"]["id"], "amount_idr": amount, "kind": "DP", "method": method,
        "received_at": timezone.now().isoformat(), "reference": f"REF-{uuid4().hex[:10]}",
        "allocations": [{"booking_participant_id": line["participant_id"], "amount_idr": amount}],
    }
    if bank:
        body["bank_account"] = bank
    payment = ok(post(travel["client"], "/api/v1/payments", body))
    return ok(post(client or travel["client"], f"/api/v1/payments/{payment['id']}/verify"))


def expense(client, **body):
    payload = {"spent_on": TODAY.isoformat(), "category": "HOTEL", "amount_idr": 1_000, "description": "Biaya uji", **body}
    return post(client, "/api/v1/expenses", payload)


@pytest.fixture
def books(billing):
    """Produksi: rekening, DP transfer 10 jt + tunai 5 jt, DP 1 jt direversal, 3 pengeluaran (1 batal), wallet 100 rb."""
    s = billing
    travel = s["travel"]
    client = travel["client"]
    bank = ok(post(client, "/api/v1/bank-accounts", {"bank_name": "BSI", "account_name": "PT Travel A", "account_number": "7012345678"}))["bank_account"]
    fund(s, 100_000)
    pay(travel, 0, 10_000_000, method="TRANSFER", bank=bank["id"])
    pay(travel, 1, 5_000_000)
    reversed_payment = pay(travel, 2, 1_000_000)
    ok(post(staff(travel["tenant"], "Finance"), f"/api/v1/payments/{reversed_payment['id']}/reverse", {"reason": "Transfer ganda"}))
    hotel = ok(expense(client, amount_idr=3_000_000, departure=travel["departure"]["id"], bank_account=bank["id"], description="Hotel Makkah"))
    ok(expense(client, amount_idr=500_000, category="OFFICE", description="Listrik kantor"))
    wrong = ok(expense(client, amount_idr=999_000, category="OTHER", description="Salah input"))
    ok(post(client, f"/api/v1/expenses/{wrong['id']}/void", {"reason": "Salah input nominal"}))
    return {**s, "bank": bank, "hotel": hotel, "client": client, "travel": travel}


def test_rpt01_finance_report_equals_source_records(books):
    report = ok(books["client"].get(f"/api/v1/reports/finance{WIDE}"))
    values = kpis(report)
    invoices = TravelInvoice.objects.without_tenant_guard().exclude(status="VOIDED")
    payments = TravelPayment.objects.without_tenant_guard()
    expenses = TravelExpense.objects.without_tenant_guard().filter(voided_at__isnull=True)
    assert values["invoiced"] == invoices.aggregate(t=Sum("total_idr"))["t"] == 88_000_000
    gross = payments.filter(status__in=["VERIFIED", "REVERSED"]).aggregate(t=Sum("amount_idr"))["t"]
    reversed_total = payments.filter(status="REVERSED").aggregate(t=Sum("amount_idr"))["t"]
    assert values["received"] == gross - reversed_total == 15_000_000
    assert values["outstanding"] == invoices.aggregate(t=Sum("balance_idr"))["t"] == 73_000_000
    assert values["expenses"] == expenses.aggregate(t=Sum("amount_idr"))["t"] == 3_500_000
    assert values["cash_net"] == 15_000_000 - 3_500_000
    assert all(check["ok"] for check in report["checks"]), report["checks"]
    cash = {row["account"]: row for row in section(report, "cash")["rows"]}
    assert cash["BSI ******5678"] == {"account": "BSI ******5678", "opening_idr": 0, "inflow_idr": 10_000_000, "reversal_idr": 0, "expense_idr": 3_000_000, "transfer_in_idr": 0, "transfer_out_idr": 0, "closing_idr": 7_000_000}
    assert cash["Tunai"]["inflow_idr"] == 6_000_000 and cash["Tunai"]["reversal_idr"] == 1_000_000 and cash["Tunai"]["closing_idr"] == 4_500_000
    receipts = {row["kind"]: row["amount_idr"] for row in section(report, "receipts")["rows"]}
    assert receipts == {"DP": 16_000_000, "Reversal": -1_000_000}
    [receivable] = section(report, "receivables")["rows"]
    assert receivable["balance_idr"] == 73_000_000 and receivable["aging"] == "Belum jatuh tempo"
    categories = {row["category"]: row["amount_idr"] for row in section(report, "expense_categories")["rows"]}
    assert categories == {"Hotel": 3_000_000, "Operasional kantor": 500_000}
    [margin] = section(report, "margin")["rows"]
    assert margin["sales_idr"] == 88_000_000 and margin["expense_idr"] == 3_000_000 and margin["margin_idr"] == 85_000_000
    assert margin["received_idr"] == 15_000_000 and margin["note"] == "Estimasi"
    current_month = ok(books["client"].get("/api/v1/reports/finance"))  # berangkat bulan depan, invoice & biaya bulan ini
    assert [row["sales_idr"] for row in section(current_month, "margin")["rows"]] == [88_000_000]


def test_cash_opening_balance_carries_prior_movements(books):
    start = (TODAY + timedelta(days=1)).isoformat()
    report = ok(books["client"].get(f"/api/v1/reports/finance?from={start}&to={(TODAY + timedelta(days=30)).isoformat()}"))
    cash = {row["account"]: row for row in section(report, "cash")["rows"]}
    assert cash["BSI ******5678"]["opening_idr"] == 7_000_000 and cash["BSI ******5678"]["inflow_idr"] == 0
    assert cash["Tunai"]["opening_idr"] == 4_500_000 and cash["Tunai"]["closing_idr"] == 4_500_000
    assert kpis(report)["received"] == 0 and kpis(report)["outstanding"] == 73_000_000  # piutang = per saat laporan


def test_rpt02_saas_wallet_report_is_separate_from_travel_cash(books):
    saas = ok(books["client"].get(f"/api/v1/reports/saas{WIDE}"))
    values = kpis(saas)
    assert values == {"balance": 0, "topups": 100_000, "charges": 100_000, "pending": 0, "refunds": 0}
    statement = section(saas, "statement")
    assert [row["amount_idr"] for row in statement["rows"]] == [0, 100_000, 0, -100_000] and statement["totals"]["amount_idr"] == 0
    ledger = WalletLedger.objects.without_tenant_guard().order_by("-sequence").first()
    assert statement["totals"]["amount_idr"] == ledger.balance_after_idr
    assert all(check["ok"] for check in saas["checks"]) and "terpisah" in saas["note"]
    assert len(section(saas, "charges")["rows"]) == 2
    finance = ok(books["client"].get(f"/api/v1/reports/finance{WIDE}"))
    cash_total = section(finance, "cash")["totals"]
    assert cash_total["inflow_idr"] == 16_000_000  # top-up wallet 100 rb tidak masuk kas travel
    assert "SaaS Wallet" in finance["note"]


def test_operational_report_readiness_occupancy_and_filters(books):
    client, travel = books["client"], books["travel"]
    report = ok(client.get(f"/api/v1/reports/operational{WIDE}"))
    [row] = section(report, "departures")["rows"]
    assert row["code"] == travel["departure"]["code"] and row["capacity"] == 10 and row["seats_taken"] == 3
    assert row["occupancy"] == 30.0 and row["participants"] == 3 and row["bookings"] == 1
    assert row["dp_received"] == 2 and row["manifest"] == "Belum" and row["readiness"] == "Lengkap"
    assert kpis(report)["occupancy"] == 30.0 and all(check["ok"] for check in report["checks"])
    [package] = section(report, "packages")["rows"]
    assert package["participants"] == 3 and package["departures"] == 1
    narrow = ok(client.get(f"/api/v1/reports/operational{WIDE}&package={travel['package']['id']}&status=OPEN"))
    assert len(section(narrow, "departures")["rows"]) == 1 and narrow["filter_labels"]["package"] == "Umroh Contoh"
    assert section(ok(client.get(f"/api/v1/reports/operational{WIDE}&status=READY")), "departures")["rows"] == []
    assert section(ok(client.get("/api/v1/reports/operational")), "departures")["rows"] == []  # bawaan bulan berjalan
    for bad in [f"?package={uuid4()}", "?status=BOGUS", "?from=2026-10-10&to=2026-10-01", "?from=10-10-2026", "?from=2020-01-01&to=2026-01-01", "?departure=x"]:
        assert client.get(f"/api/v1/reports/operational{bad}").status_code == 400, bad
    assert client.get("/api/v1/reports/unknown").status_code == 404


def test_exports_pdf_xlsx_csv_are_audited(books):
    client = books["client"]
    pdf = client.get(f"/api/v1/reports/finance{WIDE}&file=pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF-") and pdf["Content-Disposition"].startswith('attachment; filename="laporan-finance-')
    assert "no-store" in pdf["Cache-Control"]
    xlsx = client.get(f"/api/v1/reports/finance{WIDE}&file=xlsx")
    book = load_workbook(BytesIO(xlsx.content))
    assert book.sheetnames[0] == "Ringkasan" and "Kas & bank per rekening" in book.sheetnames
    cash = book["Kas & bank per rekening"]
    assert cash["A1"].value == "Rekening" and cash.cell(cash.max_row, 8).value == 11_500_000
    csv = client.get(f"/api/v1/reports/operational{WIDE}&file=csv")
    assert csv.status_code == 200 and "Kesiapan per keberangkatan" in csv.content.decode("utf-8-sig")
    assert client.get(f"/api/v1/reports/saas{WIDE}&file=pdf").content.startswith(b"%PDF-")
    assert client.get(f"/api/v1/reports/finance{WIDE}&file=doc").status_code == 400
    assert AuditEvent.objects.filter(action="reports.exported").count() == 4
    assert export._csv_value("=HYPERLINK(1)") == "'=HYPERLINK(1)" and export._csv_value(-5) == -5


def test_report_and_expense_permissions(books):
    tenant = books["travel"]["tenant"]
    operations, finance, sales = staff(tenant, "Operations"), staff(tenant, "Finance"), staff(tenant, "Sales")
    assert operations.get("/api/v1/reports/operational").status_code == 200
    assert operations.get("/api/v1/reports/finance").status_code == 403
    assert operations.get("/api/v1/reports/saas").status_code == 403
    assert operations.get("/api/v1/expenses").status_code == 403
    for kind in ["operational", "finance", "saas"]:
        assert finance.get(f"/api/v1/reports/{kind}").status_code == 200
        assert sales.get(f"/api/v1/reports/{kind}").status_code == 403
    assert expense(finance, amount_idr=1).status_code == 201
    assert expense(sales).status_code == 403


def test_expense_validation_void_and_immutability(books):
    client = books["client"]
    assert expense(client, spent_on=(TODAY + timedelta(days=2)).isoformat()).status_code == 400
    assert expense(client, departure=str(uuid4())).status_code == 400
    assert expense(client, bank_account=str(uuid4())).status_code == 400
    assert expense(client, amount_idr=0).status_code == 400
    assert expense(client, category="PESTA").status_code == 400
    hotel = books["hotel"]
    assert post(client, f"/api/v1/expenses/{hotel['id']}/void", {"reason": "x"}).status_code == 400
    voided = ok(post(client, f"/api/v1/expenses/{hotel['id']}/void", {"reason": "Hotel dibatalkan vendor"}))
    again = ok(post(client, f"/api/v1/expenses/{hotel['id']}/void", {"reason": "Alasan lain sekali"}))
    assert voided["status"] == again["status"] == "VOIDED" and again["void_reason"] == "Hotel dibatalkan vendor"
    listing = ok(client.get("/api/v1/expenses?status=RECORDED"))
    assert listing["count"] == 1 and listing["total_idr"] == 500_000
    assert ok(client.get(f"/api/v1/expenses?departure={books['travel']['departure']['id']}"))["count"] == 1
    assert client.get("/api/v1/expenses?from=bad").status_code == 400
    row = TravelExpense.objects.without_tenant_guard().get(pk=hotel["id"])
    row.amount_idr = 1
    with pytest.raises(DjangoValidationError):
        row.save()
    with pytest.raises(DjangoValidationError):
        row.delete()
    assert AuditEvent.objects.filter(action="expenses.voided").count() == 2  # salah input + hotel


def test_tenant_isolation_and_rls(books, tenant_b):
    other = staff(tenant_b, "Owner")
    from modules.tenants.models import TenantEnvironment

    TenantEnvironment.objects.create(tenant=tenant_b, environment="production", state="ACTIVE")
    values = kpis(ok(other.get(f"/api/v1/reports/finance{WIDE}")))
    assert values == {"invoiced": 0, "received": 0, "outstanding": 0, "expenses": 0, "cash_net": 0}
    assert ok(other.get("/api/v1/expenses"))["count"] == 0
    assert post(other, f"/api/v1/expenses/{books['hotel']['id']}/void", {"reason": "bukan milik saya"}).status_code == 404
    with connection.cursor() as cursor:
        cursor.execute("SELECT relrowsecurity, relforcerowsecurity FROM pg_class WHERE relname='travel_finance_travelexpense'")
        assert cursor.fetchone() == (True, True)


def test_demo_reports_are_marked(scenario):
    s = scenario(environment="demo")
    report = ok(s["client"].get(f"/api/v1/reports/finance{WIDE}"))
    assert report["environment"] == "demo"
    assert s["client"].get(f"/api/v1/reports/finance{WIDE}&file=pdf")["Content-Disposition"].startswith('attachment; filename="DEMO-laporan-finance-')
    saas = ok(s["client"].get("/api/v1/reports/saas"))
    assert saas["kpis"] == [] and "demo" in saas["note"].lower()
