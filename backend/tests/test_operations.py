"""S15: operasional keberangkatan — status pembayaran manifest, kesiapan, ekspor PDF/XLSX, cetak MOD-12."""
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4

import pytest
from django.core.files.storage import default_storage
from django.utils import timezone
from openpyxl import load_workbook

from modules.audit.models import AuditEvent
from modules.departures.models import Departure
from modules.manifest.models import ManifestVersion

from .test_payment_security import staff
from .test_payments import ok, post
from .test_payments import scenario as scenario


@pytest.fixture(autouse=True)
def private_media(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path


def pay(s, index, amount):
    line = s["invoice"]["lines"][index]
    payment = ok(post(s["client"], "/api/v1/payments", {
        "invoice": s["invoice"]["id"], "amount_idr": amount, "kind": "DP", "method": "CASH",
        "received_at": timezone.now().isoformat(), "reference": f"REF-{uuid4().hex[:10]}",
        "allocations": [{"booking_participant_id": line["participant_id"], "amount_idr": amount}],
        "simulate": s["environment"] == "demo",
    }))
    ok(post(s["client"], f"/api/v1/payments/{payment['id']}/verify"))
    return line


def manifest(s, query=""):
    return ok(s["client"].get(f"/api/v1/departures/{s['departure']['id']}/manifest{query}"))


@pytest.fixture
def ops(scenario):
    s = scenario(environment="production")
    full = pay(s, 0, s["invoice"]["lines"][0]["amount_idr"])
    partial = pay(s, 1, 1_000_000)
    s["full"], s["partial"] = full, partial
    return s


def test_manifest_rows_carry_payment_status_and_filters(ops):
    rows = {r["participant_id"]: r for r in manifest(ops)["rows"]}
    full, partial = rows[ops["full"]["participant_id"]], rows[ops["partial"]["participant_id"]]
    assert full["payment_status"] == "paid" and full["balance_idr"] == 0 and full["dp_received"]
    assert partial["payment_status"] == "partial" and partial["paid_idr"] == 1_000_000
    assert partial["balance_idr"] == partial["amount_idr"] - 1_000_000 and partial["room_type"] == "TRIPLE"
    assert [r["payment_status"] for r in manifest(ops, "?payment=unpaid")["rows"]] == ["unpaid"]
    assert len(manifest(ops, f"?booking={ops['booking']['id']}")["rows"]) == 3
    assert manifest(ops, f"?booking={uuid4()}")["rows"] == []
    for bad in ["?payment=lunas", "?booking=bukan-uuid"]:
        assert ops["client"].get(f"/api/v1/departures/{ops['departure']['id']}/manifest{bad}").status_code == 400


def test_readiness_summary_ignores_filters(ops):
    ready = manifest(ops, "?payment=paid")["readiness"]
    assert ready["participants"] == 3 and ready["dp_received"] == 2 and ready["fully_paid"] == 1
    assert ready["outstanding_idr"] == sum(line["amount_idr"] for line in ops["invoice"]["lines"]) - ops["full"]["amount_idr"] - 1_000_000
    assert ready["checklist_ready"] and ready["gap_count"] == 0  # no requirements configured yet
    ok(post(ops["client"], "/api/v1/document-requirements", {"package": ops["package"]["id"], "kind": "passport", "mandatory": True, "deadline_rule": "before_ready"}))
    ready = manifest(ops)["readiness"]
    assert not ready["checklist_ready"] and ready["gap_count"] == 3 and ready["docs_complete"] == 0


def test_xlsx_export_masked_by_default_and_formula_safe(ops):
    pilgrim = ops["pilgrims"][2]
    from .test_payments import patch as http_patch

    ok(http_patch(ops["client"], f"/api/v1/pilgrims/{pilgrim['id']}", {"name": "=HYPERLINK(\"http://evil\")", "identity_number": "3175000000000099"}))
    response = ops["client"].get(f"/api/v1/departures/{ops['departure']['id']}/manifest/export?file=xlsx")
    assert response.status_code == 200 and response["Content-Type"].startswith("application/vnd.openxmlformats")
    assert response["Content-Disposition"].endswith('.xlsx"') and "no-store" in response["Cache-Control"]
    sheet = load_workbook(BytesIO(response.content)).active
    assert sheet["A5"].value == "No" and sheet["L5"].value == "Pembayaran"
    cells = [cell for row in sheet.iter_rows(min_row=6) for cell in row]
    evil = next(cell for cell in cells if isinstance(cell.value, str) and cell.value.startswith("=HYPERLINK"))
    assert evil.data_type == "s"  # stays text, never a formula
    values = " ".join(str(cell.value) for cell in cells)
    assert "3175000000000099" not in values and "0099" in values
    event = AuditEvent.objects.filter(action="manifest.exported").latest("occurred_at")
    assert event.changes_json["format"] == "xlsx" and event.changes_json["pii"] is False


def test_pii_exports_require_permission_and_are_audited(ops):
    url = f"/api/v1/departures/{ops['departure']['id']}/manifest/export?file=pdf&pii=1"
    operations = staff(ops["tenant"], "Operations")
    assert operations.get(url).status_code == 403
    assert AuditEvent.objects.filter(action="manifest.export_denied").count() == 1
    owner = ops["client"].get(url)
    assert owner.status_code == 200 and owner.content.startswith(b"%PDF-")
    assert operations.get(f"/api/v1/departures/{ops['departure']['id']}/manifest/export?file=pdf").status_code == 200
    assert operations.get(f"/api/v1/departures/{ops['departure']['id']}/manifest/export?file=doc").status_code == 400
    assert staff(ops["tenant"], "Sales").get(url.replace("&pii=1", "")).status_code == 403


def test_published_version_pdf_from_verified_snapshot(ops):
    ok(post(ops["client"], f"/api/v1/departures/{ops['departure']['id']}/manifest/publish"))
    base = f"/api/v1/departures/{ops['departure']['id']}/manifest/versions"
    response = ops["client"].get(f"{base}/1/pdf")
    assert response.status_code == 200 and response.content.startswith(b"%PDF-")
    assert response["Content-Disposition"].endswith(f'-{ops["departure"]["code"]}-v1.pdf"')
    assert ops["client"].get(f"{base}/9/pdf").status_code == 404
    version = ManifestVersion.objects.without_tenant_guard().get()
    with default_storage.open(version.snapshot_object_key, "wb") as handle:
        handle.write(b'{"tampered": true}')
    assert ops["client"].get(f"{base}/1/pdf").status_code == 400
    assert AuditEvent.objects.filter(action="manifest.version_printed").count() == 1


@pytest.mark.parametrize("kind", ["registration-forms", "labels", "envelopes"])
def test_print_documents(ops, kind):
    url = f"/api/v1/departures/{ops['departure']['id']}/print/{kind}"
    response = ops["client"].get(url)
    assert response.status_code == 200 and response.content.startswith(b"%PDF-")
    assert ops["client"].get(f"{url}?booking={ops['booking']['id']}").status_code == 200
    event = AuditEvent.objects.filter(action="manifest.printed").latest("occurred_at")
    assert event.changes_json["kind"] == kind and event.changes_json["rows"] == 3


def test_print_permissions_pii_and_unknown_kind(ops):
    base = f"/api/v1/departures/{ops['departure']['id']}/print"
    operations = staff(ops["tenant"], "Operations")
    assert operations.get(f"{base}/registration-forms").status_code == 200
    assert operations.get(f"{base}/registration-forms?pii=1").status_code == 403
    assert AuditEvent.objects.filter(action="manifest.print_denied").count() == 1
    assert ops["client"].get(f"{base}/registration-forms?pii=1").status_code == 200
    assert ops["client"].get(f"{base}/visa").status_code == 404
    assert staff(ops["tenant"], "Sales").get(f"{base}/labels").status_code == 403
    assert ops["client"].get(f"{base}/labels?booking=x").status_code == 400


def test_demo_documents_are_watermarked(scenario):
    s = scenario(environment="demo")
    drawn = []
    original = __import__("reportlab.pdfgen.canvas", fromlist=["Canvas"]).Canvas.drawCentredString

    def spy(self, x, y, text, *args, **kwargs):
        drawn.append(text)
        return original(self, x, y, text, *args, **kwargs)

    with patch("reportlab.pdfgen.canvas.Canvas.drawCentredString", spy):
        export = s["client"].get(f"/api/v1/departures/{s['departure']['id']}/manifest/export?file=pdf")
        labels = s["client"].get(f"/api/v1/departures/{s['departure']['id']}/print/labels")
    assert export.status_code == labels.status_code == 200
    assert export["Content-Disposition"].startswith('attachment; filename="DEMO-manifest-')
    assert labels["Content-Disposition"].startswith('attachment; filename="DEMO-labels-')
    assert drawn.count("DEMO — BUKAN DOKUMEN RESMI") >= 2


def test_other_tenant_departure_is_not_found(ops, tenant_b):
    from modules.catalog.models import TravelPackage

    package = TravelPackage.objects.without_tenant_guard().create(tenant=tenant_b, environment="production", code="PKT-OTHER", name="Lain", price_idr=1)
    foreign = Departure.objects.without_tenant_guard().create(tenant=tenant_b, environment="production", code="DPR-OTHER", package=package, start_at=timezone.now(), seat_capacity=5)
    for path in [f"/api/v1/departures/{foreign.pk}/manifest/export?file=xlsx", f"/api/v1/departures/{foreign.pk}/print/labels", f"/api/v1/departures/{foreign.pk}/manifest/versions/1/pdf"]:
        assert ops["client"].get(path).status_code == 404, path
