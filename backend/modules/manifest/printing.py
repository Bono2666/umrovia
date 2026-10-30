"""Dokumen cetak operasional (S15; PRD MOD-12, OPS-04; Spec TR05).

Manifest PDF/XLSX, formulir pendaftaran, label koper dan amplop dibuat lokal
dari data yang sudah difilter dan diizinkan pemanggil. Teks di-escape; tidak
ada aset/URL remote. Dokumen demo selalu ber-watermark "DEMO".
"""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, PageBreak, SimpleDocTemplate, Spacer, Table, TableStyle

from modules.documents import pdf_theme as theme
from modules.documents.render import format_idr
from modules.payments.pdf import document_date

GENDER = {"male": "L", "female": "P"}
PAYMENT = {"paid": "Lunas", "partial": "Sebagian", "unpaid": "Belum bayar"}
PARTICIPANT = {
    "REGISTERED": "Terdaftar", "AWAITING_DP": "Menunggu DP", "DP_VERIFIED_PENDING_CHARGE": "DP diterima",
    "CONFIRMED": "Terkonfirmasi", "READY": "Siap", "DEPARTED": "Berangkat", "COMPLETED": "Selesai",
    "CANCELED": "Batal", "REVIEW_REQUIRED": "Perlu review", "DRAFT": "Draf",
}
DOC_KIND = {"ktp": "KTP", "passport": "Paspor", "photo": "Foto", "kk": "KK", "other": "Lainnya"}
ROOM = {"QUAD": "Quad", "TRIPLE": "Triple", "DOUBLE": "Double", "SINGLE": "Single"}
DL = (220 * mm, 110 * mm)


def _mask(value):
    """Nomor tersamar dari server ("************0001") → "•••• 0001" seperti di aplikasi."""
    return value.lstrip("*").join(["•••• ", ""]) if value and value.startswith("*") else value


def _lines(values, style):
    """Baris opsional (alamat/telepon kosong tidak dicetak sebagai tanda strip)."""
    return [theme.p(value, style) for value in values if value and value != "-"]


def _decor(meta, *, footer=True):
    """Footer (judul, pencetak, halaman) + watermark DEMO untuk environment demo."""
    return theme.page_decor(footer_left=f"{meta['title']} · {meta['travel']} · dicetak {meta['printed_at']} oleh {meta['printed_by']}",
                            demo=meta.get("demo"), demo_text="DEMO — BUKAN DOKUMEN RESMI", demo_size=44, footer=footer)


def print_meta(*, title, travel, environment, actor, zone="Asia/Jakarta"):
    return {
        "title": title, "travel": travel, "demo": environment == "demo",
        "printed_at": document_date(datetime.now().astimezone().isoformat(), zone, with_time=True),
        "printed_by": (getattr(actor, "full_name", "") or getattr(actor, "email", "")) or "sistem",
    }


def _departure_cells(dep):
    when = document_date(dep["start_at"])
    if dep.get("end_at"):
        when += f" – {document_date(dep['end_at'])}"
    cells = [("Paket", dep["package_name"], dep["code"]), ("Jadwal", when), ("Kota asal", dep.get("city_of_origin") or "—")]
    if dep.get("route"):
        cells.append(("Rute", dep["route"]))
    if dep.get("meeting_point"):
        cells.append(("Titik kumpul", dep["meeting_point"]))
    return cells


# ------------------------------------------------------------- manifest --

def _manifest_row(index, row, pii):
    docs = "Lengkap" if row.get("docs_complete") else "Kurang: " + ", ".join(DOC_KIND.get(k, k) for k in row.get("docs_missing", []))
    payment = PAYMENT.get(row.get("payment_status"), "—")
    if row.get("payment_status") and row["payment_status"] != "paid":
        payment += f" · sisa {format_idr(row.get('balance_idr', 0))}"
    return [
        str(index), row["name"], GENDER.get(row.get("gender"), "—"),
        document_date(row["birth_date"]) if row.get("birth_date") else "—",
        row["booking_code"], ROOM.get(row.get("room_type"), row.get("room_type") or "—"),
        row.get("identity_number") if pii else _mask(row.get("identity_masked")),
        row.get("passport_number") if pii else _mask(row.get("passport_masked")),
        row.get("phone"), PARTICIPANT.get(row.get("participant_status"), row.get("participant_status")),
        docs, payment,
    ]


MANIFEST_HEADERS = ["No", "Nama", "L/P", "Tgl lahir", "Booking", "Kamar", "NIK", "Paspor", "Telepon", "Status", "Dokumen", "Pembayaran"]


def manifest_pdf(*, departure, rows, pii, meta, version=None) -> bytes:
    s = theme.styles()
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=28, rightMargin=28, topMargin=30, bottomMargin=42, title=f"Manifest {departure['code']}", author=meta["travel"])
    stamp = (f"Versi {version['version_no']} · diterbitkan {version['published_at']} oleh {version['published_by']}" if version
             else "Belum diterbitkan sebagai versi — salinan kerja saat dicetak")
    badges = [(f"Versi {version['version_no']}", "primary") if version else ("Salinan kerja", "warning"),
              ("Data pribadi lengkap — rahasia", "danger") if pii else ("Identitas tersamar", "neutral")]
    if meta.get("demo"):
        badges.append(("Demo", "danger"))
    story = [theme.kop(s, width=doc.width, name=meta["travel"], lines=[stamp], kind="Manifest jamaah", number=departure["code"], badges=badges, split=0.55),
             Spacer(1, 12), theme.info_grid(s, [*_departure_cells(departure), ("Jumlah jamaah", f"{len(rows)} orang")], doc.width), Spacer(1, 12)]
    # Lebar diukur dengan Poppins 7,5 pt + padding 2×4: NIK 16 digit (ekspor PII), kode booking,
    # telepon 15 digit dan "Terkonfirmasi" muat satu baris; total = lebar A4 lanskap − margin (786).
    widths = [20, 100, 22, 58, 64, 38, 78, 50, 76, 62, 108, 110]
    table = theme.data_table(s, MANIFEST_HEADERS, [_manifest_row(i, row, pii) for i, row in enumerate(rows, 1)], widths, zebra=True, hpad=4)
    story.append(table if rows else theme.notice(s, "Belum ada jamaah aktif pada keberangkatan ini.", doc.width))
    theme.build(doc, story, _decor(meta))
    return output.getvalue()


def manifest_xlsx(*, departure, rows, pii, meta) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Manifest"
    sheet.append([f"Manifest Jamaah — {meta['travel']}" + (" (DEMO)" if meta.get("demo") else "")])
    sheet.append([f"{departure['package_name']} · {departure['code']} · {document_date(departure['start_at'])} · {departure.get('city_of_origin') or ''}"])
    sheet.append([f"Dicetak {meta['printed_at']} oleh {meta['printed_by']} · " + ("data pribadi lengkap — rahasia" if pii else "identitas tersamar")])
    sheet.append([])
    headers = MANIFEST_HEADERS + ["Tagihan (Rp)", "Terbayar (Rp)", "Sisa (Rp)"]
    sheet.append(headers)
    for i, row in enumerate(rows, 1):
        values = _manifest_row(i, row, pii)
        values[0] = i
        sheet.append(values + [row.get("amount_idr", 0), row.get("paid_idr", 0), row.get("balance_idr", 0)])
    sheet["A1"].font = Font(bold=True, size=14, color="0E6B5C")
    for cell in sheet[5]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EAF3EF")
    for row in sheet.iter_rows(min_row=6):
        for cell in row:
            # Teks yang diawali =, +, -, @ dipaksa string agar tidak menjadi formula.
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for letter, width in zip("ABCDEFGHIJKLMNO", [5, 28, 5, 16, 14, 9, 20, 14, 17, 15, 24, 22, 14, 14, 14], strict=True):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A6"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


# --------------------------------------------------- formulir & label/amplop --

def registration_forms_pdf(*, departure, people, pii, meta, issuer) -> bytes:
    """Satu halaman per jamaah: data diri, kontak, perjalanan dan tanda tangan."""
    s = theme.styles()
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=42, rightMargin=42, topMargin=40, bottomMargin=50, title=f"Formulir pendaftaran {departure['code']}", author=meta["travel"])
    width = doc.width
    issuer_lines = [issuer.get("address"), issuer.get("phone") and f"Telepon {issuer['phone']}"]
    story = []
    for index, person in enumerate(people):
        if index:
            story.append(PageBreak())
        identity = [
            ("Nama sesuai paspor", person["name"]), ("Kode jamaah", person["code"]),
            ("Jenis kelamin", {"male": "Laki-laki", "female": "Perempuan"}.get(person.get("gender"), "—")),
            ("Tempat, tanggal lahir", ", ".join(filter(None, [person.get("birth_place"), document_date(person["birth_date"]) if person.get("birth_date") else ""])) or "—"),
            ("NIK", person.get("identity_number") if pii else _mask(person.get("identity_masked"))),
            ("Nomor paspor", person.get("passport_number") if pii else _mask(person.get("passport_masked"))),
            ("Masa berlaku paspor", document_date(person["passport_expires_at"]) if person.get("passport_expires_at") else "—"),
        ]
        contact = [
            ("Alamat", person.get("address")), ("Telepon / WhatsApp", person.get("phone")), ("Email", person.get("email")),
            ("Wali / mahram", " · ".join(filter(None, [person.get("guardian_name"), person.get("guardian_relation")])) or "—"),
            ("Kontak darurat", " · ".join(filter(None, [person.get("emergency_contact_name"), person.get("emergency_contact_phone")])) or "—"),
        ]
        trip = [("Booking", person["booking_code"]), ("Tipe kamar", ROOM.get(person.get("room_type"), person.get("room_type") or "—"))]
        badges = [("Data lengkap — rahasia", "danger")] if pii else []
        if meta.get("demo"):
            badges.append(("Demo", "danger"))
        story += [
            theme.kop(s, width=width, name=issuer["legal_name"], lines=issuer_lines, kind="Formulir pendaftaran jamaah", number=person["code"], badges=badges),
            Spacer(1, 12), theme.info_grid(s, _departure_cells(departure), width, cols=3), Spacer(1, 12),
            theme.form_section(s, "Data diri", identity, width), Spacer(1, 10),
            theme.form_section(s, "Kontak & wali", contact, width), Spacer(1, 10),
            theme.form_section(s, "Perjalanan", trip, width), Spacer(1, 12),
            theme.p("Dengan menandatangani formulir ini, jamaah menyatakan data di atas benar dan menyetujui ketentuan perjalanan dari travel.", s["small"]),
            Spacer(1, 12), KeepTogether(_signatures(s, width, person["name"])),
        ]
    if not people:
        story.append(theme.notice(s, "Belum ada jamaah aktif pada keberangkatan ini.", width))
    theme.build(doc, story, _decor(meta))
    return output.getvalue()


def _signatures(s, width, name):
    def box(title, signer):
        return [theme.p(title, s["label"]), Spacer(1, 48), theme.p(signer, s["body"]), theme.p("Tanggal: ....................................", s["small"])]

    half = (width - 14) / 2
    table = Table([[box("Tanda tangan jamaah", f"( {name} )"), "", box("Petugas travel", "( ............................................ )")]], colWidths=[half, 14, half])
    table.setStyle(TableStyle(theme.squared([
        ("BOX", (0, 0), (0, 0), 0.75, theme.LINE_STRONG), ("BOX", (2, 0), (2, 0), 0.75, theme.LINE_STRONG),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 10), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
    ])))
    return table


def _framed(content, width, height=None, padding=8):
    """Blok berbingkai dengan aksen garis kiri hijau (label koper, penerima amplop)."""
    table = Table([[content]], colWidths=[width], rowHeights=[height] if height else None)
    table.setStyle(TableStyle(theme.squared([
        ("BOX", (0, 0), (-1, -1), 0.75, theme.LINE_STRONG), ("LINEBEFORE", (0, 0), (0, -1), 3, theme.PRIMARY),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), padding + 4), ("RIGHTPADDING", (0, 0), (-1, -1), padding),
        ("TOPPADDING", (0, 0), (-1, -1), padding - 2), ("BOTTOMPADDING", (0, 0), (-1, -1), padding),
    ])))
    return table


def labels_pdf(*, departure, people, meta) -> bytes:
    """Label koper/nama A4: 2 kolom × 7 baris (≈ 95 × 36 mm), tiap label berbingkai."""
    s = theme.styles()
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=14 * mm, title=f"Label jamaah {departure['code']}", author=meta["travel"])
    cells = []
    for person in people:
        cells.append(_framed([
            theme.p(meta["travel"], theme.tone(s, "label", theme.PRIMARY)), Spacer(1, 2),
            theme.p(person["name"], s["name"]),
            theme.p(f"{person['booking_code']} · {ROOM.get(person.get('room_type'), person.get('room_type') or '')}".strip(" ·"), s["cell"]),
            theme.p(f"{departure['package_name']} · {document_date(departure['start_at'])}", s["small"]),
            *_lines([person.get("phone")], s["small"]),
        ], 91 * mm, 32 * mm))
    if len(cells) % 2:
        cells.append("")
    grid = [cells[i:i + 2] for i in range(0, len(cells), 2)]
    story = [theme.notice(s, "Belum ada jamaah aktif pada keberangkatan ini.", doc.width)]
    if grid:
        table = Table(grid, colWidths=[95 * mm, 95 * mm], rowHeights=[36 * mm] * len(grid))
        table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                                   ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
        story = [table]
    theme.build(doc, story, _decor(meta))
    return output.getvalue()


def envelopes_pdf(*, people, meta, issuer) -> bytes:
    """Amplop DL (220 × 110 mm) per jamaah: pengirim travel, penerima jamaah dalam bingkai."""
    s = theme.styles()
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=DL, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=11 * mm, bottomMargin=10 * mm, title="Amplop jamaah", author=meta["travel"])
    story = []
    for index, person in enumerate(people):
        if index:
            story.append(PageBreak())
        sender = [theme.p("Pengirim", theme.tone(s, "label", theme.PRIMARY)), Spacer(1, 2), theme.p(issuer["legal_name"], s["bold"]),
                  *_lines([issuer.get("address"), issuer.get("phone")], s["small"])]
        recipient = _framed([theme.p("Kepada Yth.", s["label"]), Spacer(1, 3), theme.p(person["name"], s["name"]),
                             theme.p(person.get("address") or "Alamat belum diisi", s["body"]), *_lines([person.get("phone")], s["small"])], 112 * mm, padding=10)
        story += [Table([[sender, ""]], colWidths=[110 * mm, 86 * mm], style=[("LEFTPADDING", (0, 0), (-1, -1), 0)]), Spacer(1, 9 * mm),
                  Table([["", recipient]], colWidths=[84 * mm, 112 * mm], style=[("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)])]
    if not people:
        story.append(theme.notice(s, "Belum ada jamaah aktif pada keberangkatan ini.", doc.width))
    theme.build(doc, story, _decor(meta, footer=False))
    return output.getvalue()
