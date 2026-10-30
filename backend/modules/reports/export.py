"""Ekspor laporan (RPT-03/06): PDF lanskap, XLSX bernomor dan CSV dari payload yang sama dengan layar."""
import csv
from datetime import date
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle

from modules.documents import pdf_theme as theme
from modules.documents.render import format_idr
from modules.payments.pdf import document_date

MONTHS = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]


def format_value(value, kind):
    if value is None or value == "":
        return "—"
    if kind == "money":
        return ("−" if value < 0 else "") + format_idr(abs(int(value)))
    if kind == "int":
        return f"{int(value):,}".replace(",", ".")
    if kind == "percent":
        return f"{value:.1f}%".replace(".", ",")
    if kind == "date":
        day = date.fromisoformat(str(value)[:10])
        return f"{day.day} {MONTHS[day.month - 1][:3]} {day.year}"
    return str(value)


def period_text(report):
    start, end = (date.fromisoformat(report["period"][key]) for key in ("from", "to"))
    return f"{start.day} {MONTHS[start.month - 1]} {start.year} – {end.day} {MONTHS[end.month - 1]} {end.year}"


def header_lines(report, labels):
    lines = [f"{report['travel']} · Periode {period_text(report)}", f"Per {document_date(report['as_of'], report['timezone'], with_time=True)}"]
    active = [f"{name}: {labels[key]}" for key, name in (("package", "Paket"), ("departure", "Keberangkatan"), ("status", "Status"), ("agent", "Agen"), ("account", "Rekening"), ("branch", "Cabang")) if report["filters"].get(key)]
    if active:
        lines.append("Filter — " + " · ".join(active))
    if report.get("note"):
        lines.append(report["note"])
    return lines


def filename(report, ext):
    prefix = "DEMO-" if report["environment"] == "demo" else ""
    return f"{prefix}laporan-{report['report']}-{report['period']['from']}_{report['period']['to']}.{ext}"


def to_pdf(report, meta, labels) -> bytes:
    s = theme.styles()
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=28, rightMargin=28, topMargin=30, bottomMargin=42, title=report["title"], author=report["travel"])
    width = doc.width
    lines = [f"Periode {period_text(report)}", *header_lines(report, labels)[1:]]  # nama travel sudah di kop
    badges = [("Demo", "danger")] if meta.get("demo") else []
    story = [theme.kop(s, width=width, name=report["travel"], lines=lines[:2], kind="Laporan", number=report["title"], badges=badges, split=0.55), Spacer(1, 12)]
    if lines[2:]:
        story += [theme.notice(s, "\n".join(lines[2:]), width), Spacer(1, 12)]
    if report["kpis"]:
        story += [_kpi_cards(report["kpis"], s, width), Spacer(1, 16)]
    for section in report["sections"]:
        story.append(theme.p(section["title"], s["h2"]))
        if section.get("description"):
            story.append(theme.p(section["description"], s["small"]))
        story.append(Spacer(1, 6))
        columns = section["columns"]
        if not section["rows"]:
            story += [theme.notice(s, "Tidak ada data pada periode ini.", width), Spacer(1, 16)]
            continue
        numeric = {i for i, col in enumerate(columns) if col["type"] in {"money", "int", "percent"}}
        rows = [[format_value(row.get(col["key"]), col["type"]) for col in columns] for row in section["rows"]]
        totals = section.get("totals")
        if totals:
            totals = [format_value(totals.get(col["key"]), col["type"]) if col["key"] in totals else " " for col in columns]
            totals[0] = totals[0] if totals[0].strip() else "Total"
        story.append(theme.data_table(s, [col["label"] for col in columns], rows, [width / len(columns)] * len(columns), numeric=numeric, totals=totals, zebra=True))
        if section.get("truncated"):
            story += [Spacer(1, 4), theme.p("Daftar dipotong pada 500 baris; persempit filter untuk rincian lengkap.", s["small"])]
        story.append(Spacer(1, 16))
    if report.get("checks"):
        rows = [[theme.badge(s, "Sesuai" if check["ok"] else "Selisih", "success" if check["ok"] else "danger"), check["label"], check["detail"]] for check in report["checks"]]
        story += [theme.p("Rekonsiliasi terhadap catatan sumber", s["h2"]), Spacer(1, 6),
                  theme.data_table(s, ["Hasil", "Pemeriksaan", "Rincian"], rows, [70, (width - 70) * 0.45, (width - 70) * 0.55])]
    theme.build(doc, story, _decor(meta))
    return output.getvalue()


def _decor(meta):
    return theme.page_decor(footer_left=f"{meta['title']} · {meta['travel']} · dicetak {meta['printed_at']} oleh {meta['printed_by']}",
                            demo=meta.get("demo"), demo_text="DEMO — BUKAN DOKUMEN RESMI", demo_size=44)


def _kpi_cards(kpis, s, width, gap=10):
    """Kartu KPI berbingkai berjarak, seperti kartu ringkasan di layar Laporan."""
    card = (width - gap * (len(kpis) - 1)) / len(kpis)
    cells, widths, commands = [], [], [("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
                                       ("LEFTPADDING", (0, 0), (-1, -1), 10), ("RIGHTPADDING", (0, 0), (-1, -1), 10)]
    for index, kpi in enumerate(kpis):
        if index:
            cells.append("")
            widths.append(gap)
        column = len(cells)
        cells.append([theme.p(kpi["label"], s["small"]), Spacer(1, 3), theme.p(format_value(kpi["value"], kpi["type"]), s["kpi"])])
        widths.append(card)
        commands += [("BOX", (column, 0), (column, 0), 0.75, theme.LINE), ("LINEABOVE", (column, 0), (column, 0), 2, theme.PRIMARY)]
    table = Table([cells], colWidths=widths)
    table.setStyle(TableStyle(theme.squared(commands)))
    return table


def to_xlsx(report, meta, labels) -> bytes:
    book = Workbook()
    summary = book.active
    summary.title = "Ringkasan"
    summary.append([report["title"] + (" (DEMO)" if meta.get("demo") else "")])
    for line in header_lines(report, labels):
        summary.append([line])
    summary.append([f"Dicetak {meta['printed_at']} oleh {meta['printed_by']}"])
    summary.append([])
    summary.append(["Indikator", "Nilai"])
    header_row = summary.max_row
    for kpi in report["kpis"]:
        summary.append([kpi["label"], kpi["value"]])
        if kpi["type"] == "money":
            summary.cell(summary.max_row, 2).number_format = "#,##0"
    if report.get("checks"):
        summary.append([])
        summary.append(["Rekonsiliasi", "Hasil", "Rincian"])
        for check in report["checks"]:
            summary.append([check["label"], "Sesuai" if check["ok"] else "Selisih", check["detail"]])
    summary["A1"].font = Font(bold=True, size=14, color="0E6B5C")
    for cell in summary[header_row]:
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 58
    summary.column_dimensions["B"].width = 20
    summary.column_dimensions["C"].width = 60
    used = set()
    for section in report["sections"]:
        name = section["title"][:31].replace("/", "-")
        while name in used:
            name = name[:29] + str(len(used))
        used.add(name)
        sheet = book.create_sheet(name)
        columns = section["columns"]
        sheet.append([col["label"] for col in columns])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EAF3EF")
        rows = list(section["rows"])
        if section.get("totals"):
            rows.append(section["totals"])
        for row in rows:
            sheet.append([_xlsx_value(row.get(col["key"]), col["type"]) for col in columns])
            for i, col in enumerate(columns, 1):
                target = sheet.cell(sheet.max_row, i)
                if isinstance(target.value, str):
                    target.data_type = "s"  # teks diawali =, +, -, @ tidak menjadi formula
                elif col["type"] == "money":
                    target.number_format = "#,##0"
                target.alignment = Alignment(vertical="top", wrap_text=True)
        if section.get("totals"):
            for cell in sheet[sheet.max_row]:
                cell.font = Font(bold=True)
        for i, col in enumerate(columns, 1):
            sheet.column_dimensions[sheet.cell(1, i).column_letter].width = 16 if col["type"] in {"money", "int", "percent", "date"} else 30
        sheet.freeze_panes = "A2"
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _xlsx_value(value, kind):
    if value is None:
        return None
    if kind in {"money", "int"}:
        return int(value)
    if kind == "percent":
        return float(value)
    return str(value)


def to_csv(report, meta, labels) -> bytes:
    buffer = StringIO()
    writer = csv.writer(buffer)
    writer.writerow([report["title"]])
    for line in header_lines(report, labels):
        writer.writerow([line])
    for kpi in report["kpis"]:
        writer.writerow([kpi["label"], kpi["value"]])
    for section in report["sections"]:
        writer.writerow([])
        writer.writerow([section["title"]])
        writer.writerow([col["label"] for col in section["columns"]])
        rows = list(section["rows"]) + ([section["totals"]] if section.get("totals") else [])
        for row in rows:
            writer.writerow([_csv_value(row.get(col["key"])) for col in section["columns"]])
    for check in report.get("checks", []):
        writer.writerow(["Rekonsiliasi", check["label"], "Sesuai" if check["ok"] else "Selisih", check["detail"]])
    # BOM agar Excel membaca UTF-8; nilai teks berawalan formula diberi tanda kutip.
    return ("﻿" + buffer.getvalue()).encode()


def _csv_value(value):
    if value is None:
        return ""
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value
